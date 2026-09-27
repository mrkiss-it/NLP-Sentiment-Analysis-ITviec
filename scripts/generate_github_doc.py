# -*- coding: utf-8 -*-
"""Tạo file Word mô tả đầy đủ đồ án và link GitHub cho Nhóm 12."""

from pathlib import Path
import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

REPO_ROOT = Path(__file__).resolve().parents[1]
TARGET_DIR = (
    REPO_ROOT.parent
    / "Nhóm 12 - Phân tích cảm xúc đánh giá trên ITViec"
    / "Nhóm 12 - Phân tích cảm xúc đánh giá trên ITViec - Báo cáo"
)
DOC_PATH = TARGET_DIR / "Link Github và Mô tả đồ án.docx"

doc = docx.Document()

# Page setup: A4, Margins 2.5cm / 2cm / 2cm / 2cm
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.5)
sec.right_margin = Cm(2.0)

# Colors
NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
DARK = RGBColor(0x26, 0x26, 0x26)
GRAY = RGBColor(0x59, 0x59, 0x59)


def style_run(run, font="Times New Roman", size=13, bold=False, italic=False, color=DARK):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_p(text="", bold=False, italic=False, size=13, color=DARK, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=4, line_spacing=1.25):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        style_run(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading(text, level=1):
    if level == 1:
        p = add_p(text, bold=True, size=14.5, color=NAVY, space_before=14, space_after=5)
    elif level == 2:
        p = add_p(text, bold=True, size=13, color=BLUE, space_before=9, space_after=3)
    else:
        p = add_p(text, bold=True, italic=True, size=13, color=NAVY, space_before=6, space_after=2)
    return p


# ----------------- HEADER -----------------
add_p("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN - ĐHQG-HCM\nKHOA KHOA HỌC MÁY TÍNH", bold=True, size=11.5, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add_p("MÔN HỌC: XỬ LÝ NGÔN NGỮ TỰ NHIÊN (HK2/2026)", bold=True, size=12, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

# ----------------- TITLE -----------------
add_p("THÔNG TIN NỘP BÀI, LINK GITHUB & MÔ TẢ ĐỒ ÁN", bold=True, size=17, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=5)
add_p("Đề tài: PHÂN TÍCH CẢM XÚC ĐÁNH GIÁ NHÂN SỰ NGÀNH CÔNG NGHỆ THÔNG TIN TRÊN NỀN TẢNG ITVIEC", bold=True, size=12.5, color=DARK, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

# ----------------- GITHUB CALLOUT BOX -----------------
tbl = doc.add_table(rows=1, cols=1)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = tbl.rows[0].cells[0]
cell.width = Cm(16.5)

# Shading and border for callout
shd_xml = f'<w:shd {nsdecls("w")} w:fill="F2F6FA"/>'
cell._tc.get_or_add_tcPr().append(parse_xml(shd_xml))

border_xml = f"""
<w:tcBorders {nsdecls("w")}>
    <w:top w:val="single" w:sz="12" w:space="0" w:color="2E75B6"/>
    <w:left w:val="single" w:sz="24" w:space="0" w:color="1F3864"/>
    <w:bottom w:val="single" w:sz="12" w:space="0" w:color="2E75B6"/>
    <w:right w:val="single" w:sz="12" w:space="0" w:color="2E75B6"/>
</w:tcBorders>
"""
cell._tc.get_or_add_tcPr().append(parse_xml(border_xml))

cp = cell.paragraphs[0]
cp.paragraph_format.space_before = Pt(6)
cp.paragraph_format.space_after = Pt(4)
r1 = cp.add_run("🔗 KHO LƯU TRỮ MÃ NGUỒN (GITHUB REPOSITORY):\n")
style_run(r1, size=12, bold=True, color=NAVY)
r2 = cp.add_run("https://github.com/mrkiss-it/NLP-Sentiment-Analysis-ITviec\n")
style_run(r2, size=13.5, bold=True, color=BLUE)
r3 = cp.add_run("Trạng thái Repository: ")
style_run(r3, size=11, bold=True, color=DARK)
r4 = cp.add_run("Public Repository · 100% PR đã merge vào master · Đầy đủ README, requirements.txt, 55 automated tests passed 100%.")
style_run(r4, size=11, italic=True, color=GRAY)

add_p("", space_before=8)

# ----------------- SECTION 1: NHÓM THỰC HIỆN -----------------
add_heading("1. THÔNG TIN NHÓM THỰC HIỆN — NHÓM 12", level=1)
p_adv = add_p("• Giảng viên hướng dẫn: ", bold=True)
r_adv = p_adv.add_run("Thầy Đặng Văn Thìn (Khoa Khoa học Máy tính)")
style_run(r_adv, bold=False)

# Table of members
t_mem = doc.add_table(rows=5, cols=4)
t_mem.alignment = WD_TABLE_ALIGNMENT.CENTER
hdrs = ["STT", "Họ và tên sinh viên", "Mã số sinh viên", "Vai trò & Phân công chính"]
col_w = [Cm(1.2), Cm(5.0), Cm(3.2), Cm(7.1)]

for i, h in enumerate(hdrs):
    c = t_mem.rows[0].cells[i]
    c.width = col_w[i]
    cp = c.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cp.add_run(h)
    style_run(r, size=11, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    c._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="1F3864"/>'))

members_data = [
    ("1", "Trần Hoàng Hôn", "26410046", "Trưởng nhóm / Business Understanding, Thu thập & Tiền xử lý dữ liệu"),
    ("2", "Nguyễn Duy Khang", "26410055", "Machine Learning Modeling, Tinh chỉnh siêu tham số & Stacking Ensemble"),
    ("3", "Vũ Văn Duy", "26410031", "Feature Engineering, Thí nghiệm đối chứng (Ablation) & Báo cáo toàn văn"),
    ("4", "Phạm Thành Trung", "26410141", "Model Evaluation, Khai phá Insight doanh nghiệp & Phát triển Web Demo"),
]

for row_idx, data in enumerate(members_data, start=1):
    row = t_mem.rows[row_idx]
    for col_idx, text in enumerate(data):
        c = row.cells[col_idx]
        c.width = col_w[col_idx]
        cp = c.paragraphs[0]
        cp.paragraph_format.space_before = Pt(3)
        cp.paragraph_format.space_after = Pt(3)
        if col_idx in (0, 2):
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = cp.add_run(text)
        style_run(r, size=11, bold=(col_idx == 1 and row_idx == 1))
        if row_idx % 2 == 0:
            c._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F2F6FA"/>'))

for row in t_mem.rows:
    for cell in row.cells:
        cell_borders = f"""
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="D9D9D9"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="D9D9D9"/>
            <w:right w:val="none"/>
        </w:tcBorders>
        """
        cell._tc.get_or_add_tcPr().append(parse_xml(cell_borders))

add_p("", space_before=8)

# ----------------- SECTION 2: MÔ TẢ ĐỀ TÀI & Ý TƯỞNG GIẢI PHÁP -----------------
add_heading("2. MÔ TẢ BÀI TOÁN & PHƯƠNG PHÁP THỰC HIỆN", level=1)

add_heading("2.1. Bản chất dữ liệu & Thách thức", level=2)
add_p("• Dữ liệu bài toán: Thu thập từ nền tảng tuyển dụng ITviec, gồm 8.417 đánh giá thực tế của nhân viên và ứng viên (180 công ty công nghệ, thời gian từ 07/2016 đến 05/2025). Mỗi đánh giá gồm 3 phần: Tiêu đề (Title), Điểm thích (What I liked) và Góp ý (Suggestions for improvement).", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
add_p("• Ghép chuỗi văn bản hợp nhất: raw_text = Title + ' . ' + What I liked + ' . ' + Suggestions for improvement để mô hình có ngữ cảnh đầy đủ của toàn bộ đánh giá.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
add_p("• Nhãn cảm xúc 3 lớp từ số sao: Tích cực (4-5★: 73,76%), Trung tính (3★: 19,47%), Tiêu cực (1-2★: 6,77%). Dữ liệu có hiện tượng mất cân bằng lớp nghiêm trọng (khoảng 11 : 1).", align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_heading("2.2. Pipeline tiền xử lý tiếng Việt chuyên sâu 2 tầng (Dual-tier)", level=2)
add_p("• Tầng 1 (clean_basic_text): Chuẩn hóa Unicode dựng sẵn NFC, loại bỏ URL/email, ánh xạ Emoji/Emojicon sang từ ngữ cảm xúc tiếng Việt, chuẩn hóa teencode và thuật ngữ IT Anh-Việt (OT, layoff, micromanage, dev, leader...). Giữ nguyên cấu trúc câu tự nhiên phục vụ Transformer (ViSoBERT).", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
add_p("• Tầng 2 (clean_advance_text): Tách từ ghép tiếng Việt bằng underthesea, lọc bỏ stopword vô nghĩa nhưng bảo lưu nghiêm ngặt các từ phủ định/định lượng ('không', 'chưa', 'ít', 'chẳng', 'thiếu') để tối ưu không gian vector hóa cho Machine Learning.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
add_p("• Nhận diện phạm vi phủ định (Negation Scope Detection): Kết hợp đối sánh cụm từ tham lam, đảo ngược cực tính từ vựng khi gặp cấu trúc phủ định ('môi trường không được thân thiện', 'lương chưa tương xứng'), nâng độ bao phủ từ điển từ 12,26% lên 99,75%.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_heading("2.3. Trích xuất đặc trưng & Huấn luyện mô hình", level=2)
add_p("• Đặc trưng: TF-IDF n-gram (1-2) 5.000 chiều văn bản kết hợp 5 đặc trưng Lexicon cảm xúc.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
add_p("• Thí nghiệm Ablation Study: Chứng minh việc đưa điểm khía cạnh (aspect rating) vào gây ra hiện tượng Data Shortcut nguy hiểm (Macro F1 ảo lên 0,74 dù không đọc một từ văn bản nào). Nhóm kiên quyết loại bỏ điểm khía cạnh để xây dựng hệ thống NLP hiểu ngôn ngữ thực thụ.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
add_p("• Mô hình & Kết quả: Thực nghiệm qua 4 mô hình học máy (Multinomial NB, Logistic Regression, Linear SVM, Random Forest), mô hình kết hợp Stacking Ensemble và đối sánh với Deep Learning ViSoBERT zero-shot. Stacking Ensemble đạt Macro F1 0,5619 trên 5-fold CV và 0,5834 trên tập Test độc lập.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)

# ----------------- SECTION 3: HƯỚNG DẪN CÀI ĐẶT & CHẠY -----------------
add_heading("3. HƯỚNG DẪN CÀI ĐẶT & KHỞI CHẠY WEB DEMO (REQUIREMENTS)", level=1)

add_p("Yêu cầu môi trường: Python >= 3.10 (Khuyến nghị Python 3.10 hoặc 3.11).")

add_heading("Bước 1: Clone mã nguồn từ GitHub", level=2)
p_code1 = add_p("git clone https://github.com/mrkiss-it/NLP-Sentiment-Analysis-ITviec.git\ncd NLP-Sentiment-Analysis-ITviec", bold=True, size=11, color=NAVY, space_before=2, space_after=5)
p_code1.paragraph_format.left_indent = Cm(0.8)

add_heading("Bước 2: Cài đặt thư viện phụ thuộc (requirements.txt)", level=2)
p_code2 = add_p("pip install -r requirements.txt", bold=True, size=11, color=NAVY, space_before=2, space_after=5)
p_code2.paragraph_format.left_indent = Cm(0.8)
add_p("(Tệp requirements.txt đã bao gồm đầy đủ: scikit-learn, underthesea, streamlit, transformers, torch, wordcloud, python-pptx, python-docx, reportlab, pytest...)", italic=True, size=11, color=GRAY)

add_heading("Bước 3: Chạy bộ kiểm thử tự động (Automated Test Suite)", level=2)
p_code3 = add_p("pytest tests/ -v", bold=True, size=11, color=NAVY, space_before=2, space_after=5)
p_code3.paragraph_format.left_indent = Cm(0.8)
add_p("Kết quả: Toàn bộ 55/55 bài kiểm thử tự động passed 100%, kiểm chứng toàn diện từ tiền xử lý, trích xuất đặc trưng, mô hình đến giao diện web.", italic=True, size=11, color=GRAY)

add_heading("Bước 4: Khởi chạy Ứng dụng Web Demo (Streamlit App)", level=2)
p_code4 = add_p("python -m streamlit run app.py", bold=True, size=11, color=NAVY, space_before=2, space_after=5)
p_code4.paragraph_format.left_indent = Cm(0.8)
add_p("Ứng dụng web sẽ tự động mở tại địa chỉ http://localhost:8501 (hoặc cổng khả dụng), cung cấp 4 phân hệ chính: Tổng quan đề tài; Dashboard Insight doanh nghiệp; Phân tích cảm xúc thời gian thực (XAI bóc tách từ ngữ); và Benchmark Leaderboard so sánh mô hình.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)

# ----------------- SECTION 4: HỒ SƠ NỘP BÀI -----------------
add_heading("4. DANH MỤC HỒ SƠ NỘP BÀI CỦA NHÓM 12 TRÊN GOOGLE DRIVE", level=1)
add_p("1. Thư mục \"Báo cáo\": Gồm file Báo cáo toàn văn 6 chương chuẩn mẫu Khoa (Báo cáo đồ án.pdf, Báo cáo đồ án.docx) và file mô tả Link GitHub này.")
add_p("2. Thư mục \"Slides\": Gồm 15 slide Dark-tech phục vụ báo cáo bảo vệ 15 phút (Slide thuyết trình.pptx, Slide thuyết trình.pdf).")
add_p("3. Mã nguồn và Demo: Quản lý trực tuyến và công khai tại link GitHub: https://github.com/mrkiss-it/NLP-Sentiment-Analysis-ITviec.")

# Save file
TARGET_DIR.mkdir(parents=True, exist_ok=True)
doc.save(str(DOC_PATH))
print(f"SUCCESS: Created Word document ({DOC_PATH.stat().st_size / 1024:.0f} KB)")

# Export PDF
pdf_path = TARGET_DIR / "Link Github và Mô tả đồ án.pdf"
try:
    import win32com.client
    import pythoncom
    pythoncom.CoInitialize()
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    doc_word = word.Documents.Open(str(DOC_PATH.resolve()))
    doc_word.SaveAs(str(pdf_path.resolve()), FileFormat=17)
    doc_word.Close(False)
    word.Quit()
    pythoncom.CoUninitialize()
    print(f"SUCCESS: Exported PDF ({pdf_path.stat().st_size / 1024:.0f} KB)")
except Exception as e:
    print(f"Word PDF conversion note: {e}")

